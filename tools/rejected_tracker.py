#!/usr/bin/env python3
"""Build the daily 'V15 Rejected Tracker' Excel (open/rejected items of SD, MAS, PQN and RFI logs).

Usage:
  python3 tools/rejected_tracker.py --items items.json [--prev "V15 Rejected Tracker <date>.xlsx"] \
      --out "V15 Rejected Tracker 2026-10-12.xlsx" [--date 2026-10-12]

items.json = list of CURRENT open items, one per reference (latest revision only):
  {"log": "SD"|"MAS"|"PQN"|"RFI", "ref": full Unifier ref, "disc": "Architectural", "title": "...",
   "rev": 2, "code": "C"|"D"|"F"|"W" (W = RFI awaiting response), "by": "Consultant"|"Employer (ER)",
   "submitted": "2026-08-27", "url": "drive folder url"}
Rule for items.json: SD/MAS/PQN latest revision coded C, D or F; RFIs whose latest response is coded C, or
that are still awaiting a response.

--prev = the tracker Mohammed sent back (his Action / Assigned to / Target date / Remarks are carried over,
         matched by Reference No.). Previous rows that are no longer open are kept:
           latest rev now under review  -> Status 'Resubmitted – under review' (visible)
           latest rev now approved      -> Status 'Closed – approved' (row hidden)
         The script needs that information in --closed (JSON {ref: "under review"|"approved"|...}); refs
         missing from both items.json and --closed are kept as 'No longer open'.
Rows with Action = TERM and closed rows stay in the file but are hidden.
"""
import argparse, datetime, json, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

MEAN = {"C": "C – Not approved, resubmit", "D": "D – Incomplete, resubmit",
        "F": "F – Approved w/ comments, resubmit", "W": "Awaiting consultant response"}
LOGS = ["SD (Drawing)", "MAS", "PQN", "RFI"]
HDR = ["No.", "Log", "Reference No.", "Discipline", "Title", "Latest Rev", "Code", "Status", "Coded by",
       "Submitted (latest rev)", "Days since submitted", "Drive folder", "Action", "Assigned to",
       "Target date", "Remarks", "First listed"]
HR = 5
F = "Arial"


def read_prev(path):
    """{ref: {action, assigned, target, remarks, first, row values}} from a previous tracker."""
    if not path:
        return {}
    ws = load_workbook(path, data_only=True)["Tracker"]
    hdr = {str(c.value).strip(): i for i, c in enumerate(ws[HR], 1) if c.value}
    out = {}
    for r in range(HR + 1, ws.max_row + 1):
        ref = ws.cell(r, hdr["Reference No."]).value
        if not ref:
            continue
        g = lambda h: ws.cell(r, hdr[h]).value if h in hdr else None
        act = (str(g("Action") or "").strip())
        act = {"hold": "HOLD", "term": "TERM", "in progress": "In progress", "inprogress": "In progress"}.get(act.lower(), act)
        out[str(ref).strip()] = {"action": act or None, "assigned": g("Assigned to"), "target": g("Target date"),
                                 "remarks": g("Remarks"), "first": g("First listed"), "log": g("Log"),
                                 "disc": g("Discipline"), "title": g("Title"), "rev": g("Latest Rev"),
                                 "code": g("Code"), "by": g("Coded by"), "submitted": g("Submitted (latest rev)"),
                                 "url": ws.cell(r, hdr["Drive folder"]).hyperlink.target if ws.cell(r, hdr["Drive folder"]).hyperlink else None}
    return out


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, str) and re.match(r"\d{4}-\d{2}-\d{2}", v):
        return datetime.date.fromisoformat(v[:10])
    return v


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--items", required=True)
    ap.add_argument("--prev")
    ap.add_argument("--closed")
    ap.add_argument("--out", required=True)
    ap.add_argument("--date")
    a = ap.parse_args()
    today = datetime.date.fromisoformat(a.date) if a.date else datetime.date.today()
    items = json.load(open(a.items, encoding="utf-8"))
    prev = read_prev(a.prev)
    closed = json.load(open(a.closed)) if a.closed else {}
    rows = []
    seen = set()
    for it in items:
        p = prev.get(it["ref"], {})
        log = "SD (Drawing)" if it["log"] == "SD" else it["log"]
        status = MEAN[it["code"]] if not (it["log"] == "RFI" and it["code"] == "C") else "C – RFI response not accepted"
        rows.append(dict(log=log, ref=it["ref"], disc=it["disc"], title=it["title"], rev="Rev.%02d" % int(it["rev"]),
                         code=it["code"], status=status, by=it["by"], submitted=as_date(it["submitted"]), url=it.get("url"),
                         action=p.get("action"), assigned=p.get("assigned"), target=as_date(p.get("target")),
                         remarks=p.get("remarks"), first=as_date(p.get("first")) or today, hidden=p.get("action") == "TERM"))
        seen.add(it["ref"])
    for ref, p in prev.items():
        if ref in seen:
            continue
        st = closed.get(ref, "")
        status = {"under review": "Resubmitted – under review", "approved": "Closed – approved"}.get(st, "No longer open")
        rows.append(dict(log=p["log"], ref=ref, disc=p["disc"], title=p["title"], rev=p["rev"], code=p["code"], status=status,
                         by=p["by"], submitted=as_date(p["submitted"]), url=p["url"], action=p["action"], assigned=p["assigned"],
                         target=as_date(p["target"]), remarks=p["remarks"], first=as_date(p["first"]),
                         hidden=(p["action"] == "TERM" or status.startswith("Closed"))))
    order = {l: i for i, l in enumerate(LOGS)}
    rows.sort(key=lambda r: (order.get(r["log"], 9), r["ref"]))

    wb = Workbook(); ws = wb.active; ws.title = "Tracker"
    thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
    navy = PatternFill("solid", fgColor="1F3A4D"); inp = PatternFill("solid", fgColor="FFF9DB")
    ws["A1"] = "V15 – Rejected / Open Items Tracker (Shop Drawings, MAS, PQN, RFI)"
    ws["A1"].font = Font(name=F, bold=True, size=14, color="1F3A4D")
    ws["A2"] = (f"Source: Unifier logs, read {today:%d %b %Y}. Latest revision per reference coded C / D / F, plus RFIs coded C "
                "or awaiting response. Rows marked TERM and closed rows are kept but hidden (unhide rows to see them).")
    ws["A2"].font = Font(name=F, italic=True, size=9, color="595959")
    ws["A3"] = "Fill the yellow columns: Action (HOLD / TERM / In progress), Assigned to, Target date, Remarks – then send the file back."
    ws["A3"].font = Font(name=F, size=9, color="9C5700")
    for i, h in enumerate(HDR, 1):
        c = ws.cell(HR, i, h); c.font = Font(name=F, bold=True, color="FFFFFF", size=10); c.fill = navy; c.border = B
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for n, r in enumerate(rows, 1):
        R = HR + n
        vals = [n, r["log"], r["ref"], r["disc"], r["title"], r["rev"], r["code"], r["status"], r["by"], r["submitted"],
                f"=IF(J{R}=\"\",\"\",TODAY()-J{R})", None, r["action"], r["assigned"], r["target"], r["remarks"], r["first"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(R, i, v); c.font = Font(name=F, size=10); c.border = B
            c.alignment = Alignment(vertical="top", wrap_text=(i in (5, 8, 16)))
        for col in (10, 15, 17):
            ws.cell(R, col).number_format = "dd-mmm-yy"
        if r["url"]:
            c = ws.cell(R, 12, "Open folder"); c.hyperlink = r["url"]; c.font = Font(name=F, size=10, color="0563C1", underline="single")
        for i in (13, 14, 15, 16):
            ws.cell(R, i).fill = inp
        ws.cell(R, 7).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(R, 7).font = Font(name=F, size=10, bold=True, color="C00000" if r["code"] in ("C", "D", "F") else "9C5700")
        if r["hidden"]:
            ws.row_dimensions[R].hidden = True
    last = HR + max(len(rows), 1)
    dv = DataValidation(type="list", formula1='"HOLD,TERM,In progress"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Action", error="Choose HOLD, TERM or In progress")
    ws.add_data_validation(dv); dv.add(f"M{HR + 1}:M{last}")
    rng = f"A{HR + 1}:Q{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$M{HR + 1}="HOLD"'], fill=PatternFill("solid", fgColor="FCE4D6")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$M{HR + 1}="TERM"'], fill=PatternFill("solid", fgColor="E7E6E6"), font=Font(color="808080")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$M{HR + 1}="In progress"'], fill=PatternFill("solid", fgColor="DDEBF7")))
    for i, w in enumerate([5, 13, 30, 14, 60, 9, 7, 30, 13, 12, 10, 12, 13, 16, 11, 34, 11], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[HR].height = 32
    ws.freeze_panes = f"D{HR + 1}"; ws.auto_filter.ref = f"A{HR}:Q{last}"
    ws.print_title_rows = f"{HR}:{HR}"; ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True

    s = wb.create_sheet("Summary")
    s["A1"] = f"Summary by log and action – {today:%d %b %Y}"; s["A1"].font = Font(name=F, bold=True, size=13, color="1F3A4D")
    cols = ["Log", "Open (C/D/F/awaiting)", "C", "D", "F", "RFI awaiting", "HOLD", "In progress", "TERM", "No action yet"]
    for i, h in enumerate(cols, 1):
        c = s.cell(3, i, h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = navy; c.border = B
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    T = "Tracker!"
    rB, rG, rM, rH = (f"{T}${c}${HR + 1}:${c}${last}" for c in "BGMH")
    openc = f'{rH},"<>Closed*",{rH},"<>Resubmitted*",{rH},"<>No longer open"'
    for j, lg in enumerate(LOGS, 4):
        s.cell(j, 1, lg)
        s.cell(j, 2, f"=COUNTIFS({rB},$A{j},{openc})")
        for col, code in ((3, "C"), (4, "D"), (5, "F"), (6, "W")):
            s.cell(j, col, f'=COUNTIFS({rB},$A{j},{rG},"{code}",{openc})')
        for col, act in ((7, "HOLD"), (8, "In progress"), (9, "TERM")):
            s.cell(j, col, f'=COUNTIFS({rB},$A{j},{rM},"{act}",{openc})')
        s.cell(j, 10, f'=COUNTIFS({rB},$A{j},{rM},"",{openc})')
    s.cell(8, 1, "Total").font = Font(name=F, bold=True)
    for col in range(2, 11):
        L = get_column_letter(col); s.cell(8, col, f"=SUM({L}4:{L}7)").font = Font(name=F, bold=True)
    for rr in range(4, 9):
        for cc in range(1, 11):
            c = s.cell(rr, cc); c.border = B
            if rr < 8: c.font = Font(name=F)
            if cc > 1: c.alignment = Alignment(horizontal="center")
    s["A10"] = "Legend"; s["A10"].font = Font(name=F, bold=True)
    leg = [("Code C", "Not approved – revise and resubmit"), ("Code D", "Incomplete – resubmit with missing info"),
           ("Code F", "Approved with comments – resubmit"), ("RFI awaiting", "RFI still waiting for the consultant response"),
           ("HOLD", "Keep open, on hold (stays visible)"), ("In progress", "Team is working on the resubmission"),
           ("TERM", "Terminated / not needed – kept in the file, hidden"),
           ("Resubmitted", "A new revision was submitted – row stays until it is coded"),
           ("Closed", "New revision approved (A/B) – kept, hidden"),
           ("Yellow cells", "Your input: Action, Assigned to, Target date, Remarks")]
    for i, (x, y) in enumerate(leg, 11):
        s.cell(i, 1, x).font = Font(name=F, bold=True, size=10); s.cell(i, 2, y).font = Font(name=F, size=10)
    s.column_dimensions["A"].width = 16
    for L in "BCDEFGHIJ":
        s.column_dimensions[L].width = 12
    s.column_dimensions["B"].width = 14
    wb.save(a.out)
    print(json.dumps({"rows": len(rows), "hidden": sum(r["hidden"] for r in rows), "carried_actions": sum(1 for r in rows if r["action"])}))


if __name__ == "__main__":
    main()
